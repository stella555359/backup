import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

def record_all_build_test_result_to_excel(build, test_result, excel_path='build_test_results.xlsx'):
	"""
	Write build and test_result to an Excel file, automatically expanding new columns and appending rows.
	:param build: build name (string)
	:param test_result: test result (dict)
	:param excel_path: path to the Excel file, default is 'build_test_results.xlsx'
	"""
	# New row data
	row = {'build': build}
	row.update(test_result)

	# Check if the file exists
	if os.path.exists(excel_path):
		df = pd.read_excel(excel_path)
		# Add new columns if any in the new row
		for key in row.keys():
			if key not in df.columns:
				df[key] = None
		# Add missing columns in the new row for old columns
		for col in df.columns:
			if col not in row:
				row[col] = None
		# Ensure column order
		row = {col: row.get(col, None) for col in df.columns}
		# Append the new row
		df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
	else:
		df = pd.DataFrame([row])

	# Reorder columns: 'build' first, others follow in order
	cols = ['build'] + [c for c in df.columns if c != 'build']
	df = df[cols]
	# Write to Excel
	df.to_excel(excel_path, index=False, engine='openpyxl')

	# Apply formatting
	workbook = load_workbook(excel_path)
	sheet = workbook.active

	# Set header style
	header_fill = PatternFill(start_color="0000FF", end_color="0000FF", fill_type="solid")
	header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = header_font
		cell.alignment = Alignment(horizontal="left")

	# Set column widths and alignment
	sheet.column_dimensions['A'].width = 35
	for col in sheet.iter_cols(min_col=2, min_row=2):
		for cell in col:
			cell.alignment = Alignment(horizontal="left")
	for col in sheet.iter_cols(min_col=2):
		col_letter = col[0].column_letter
		sheet.column_dimensions[col_letter].width = 15

	# Set font for all other cells
	for row in sheet.iter_rows(min_row=2):
		for cell in row:
			cell.font = Font(name="Calibri", size=10)

	workbook.save(excel_path)

if __name__ == "__main__":
	# Example usage
	build1 = "SBTS26R2_ENB_0000_000281_000000"
	test_result1 = {"tput": 100, "bler": 0.01}
	record_all_build_test_result_to_excel(build1, test_result1)

	build2 = "SBTS26R2_ENB_0000_000282_000000"
	test_result2 = {"tput": 90, "bler": 0.02}
	record_all_build_test_result_to_excel(build2, test_result2)

	build3 = "SBTS26R2_ENB_0000_000283_000000"
	test_result3 = {"tput": 80, "bler": 0.03, "mumimo_pair": "98%"}
	record_all_build_test_result_to_excel(build3, test_result3)

	build4 = "SBTS26R2_ENB_0000_000284_000000"
	test_result4 = {"tput": 70}
	record_all_build_test_result_to_excel(build4, test_result4)