from __future__ import annotations

import asyncio
import html
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from playwright.async_api import async_playwright


ROOT = Path(__file__).resolve().parent
HTML_PATH = ROOT / "hainan_family_self_drive_8d7n.html"
PDF_PATH = ROOT / "hainan_family_self_drive_8d7n.pdf"
XLSX_PATH = ROOT / "hainan_family_self_drive_8d7n.xlsx"

TITLE = "海南亲子自驾 8天7晚"
SUBTITLE = "海口1晚 → 万宁1晚 → 陵水2晚 → 三亚3晚"
RESEARCHED_AT = "2026-07-21"
TRAVEL_WINDOW = "建议在 2026年8月1日至10日之间选择连续8天"
TARGET_BUDGET = 25_000


def amap(keyword: str, city: str = "海南") -> str:
    return (
        "https://uri.amap.com/search"
        f"?keyword={quote(keyword)}&city={quote(city)}&callnative=1"
    )


def safe(value: object) -> str:
    return html.escape(str(value), quote=True)


def list_html(items: list[str]) -> str:
    return "".join(f"<li>{safe(item)}</li>" for item in items)


DAYS = [
    {
        "day": "D1",
        "route": "抵达海口",
        "sleep": "海口骑楼老街 / 国贸一带",
        "drive": "约20–35公里；机场至市区约40–60分钟，以实时导航为准",
        "focus": "取车、入住、骑楼老街和海南第一餐",
        "morning": [
            "按实际航班抵达海口美兰机场，不为首日安排固定景点。",
            "取空间较大的五座SUV，根据儿童身高体重确认两套安全座椅/增高垫。",
            "现场试装后确认后排中间成人可以正常坐下并独立系安全带，行李也不遮挡视线；否则升级七座。",
        ],
        "lunch": "若中午前抵达，可在机场或市区吃简餐；避免取车后空腹长时间驾驶。",
        "afternoon": [
            "取车时录制车身、轮胎、玻璃、油量和仪表盘视频。",
            "入住带停车场的酒店，成人分工整理行李、儿童休息。",
        ],
        "evening": [
            "体力允许时逛骑楼老街1–2小时；开放式街区可灵活调整。",
            "晚餐优先尝试辣汤饭、海南粉、清补凉或老爸茶小吃。",
        ],
        "kid": "首日不安排长距离步行；给孩子留出适应气候和早睡时间。",
        "rain": "直接入住酒店，在附近商场或酒店餐厅解决晚餐；骑楼老街挪到D2早餐后。",
        "booking": "确认机场取车网点、夜间服务费、酒店停车和早餐儿童政策。",
        "nav": [
            ("海口美兰国际机场", "海口美兰国际机场", "海口"),
            ("骑楼老街", "海口骑楼老街", "海口"),
        ],
    },
    {
        "day": "D2",
        "route": "海口 → 万宁",
        "sleep": "万宁石梅湾 / 兴隆一带",
        "drive": "约150–170公里；纯驾驶约2–2.5小时，带娃按3小时预留",
        "focus": "东线转场、酒店午休、石梅湾日落",
        "morning": [
            "早餐后补逛骑楼老街或体验一顿老爸茶，控制在10:00前结束。",
            "离开海口前检查油量、饮水、防晒、儿童零食和晕车用品。",
        ],
        "lunch": "抵达兴隆或石梅湾后用午餐，推荐后安粉、海南菜或清淡简餐。",
        "afternoon": [
            "入住后先午休和玩酒店泳池，不增加途中打卡点。",
            "17:00后前往石梅湾正规开放沙滩散步、玩沙。",
        ],
        "evening": [
            "看天气决定是否等日落；雷雨前立即离开海滩。",
            "晚餐可选东山羊、和乐蟹、兴隆南洋菜；海鲜先问价再点单。",
        ],
        "kid": "海边只在救生员或酒店明确开放区域活动；不让孩子单独接近浪线。",
        "rain": "取消沙滩，使用酒店儿童乐园、泳池或兴隆镇室内餐饮；保留体力给D3植物园。",
        "booking": "优先选择有儿童泳池、室内儿童区和停车场的酒店。",
        "nav": [
            ("石梅湾", "万宁石梅湾", "万宁"),
            ("兴隆镇", "万宁兴隆镇", "万宁"),
        ],
    },
    {
        "day": "D3",
        "route": "万宁 → 陵水",
        "sleep": "陵水清水湾",
        "drive": "万宁至陵水约55–75公里；纯驾驶约1–1.5小时",
        "focus": "兴隆热带植物园、兴隆咖啡、清水湾入住",
        "morning": [
            "08:30左右进入兴隆热带植物园，安排约2小时，避开正午高温。",
            "园内以热带植物和科普为主；孩子不触摸未知果实和植物汁液。",
        ],
        "lunch": "在兴隆体验咖啡、可可或南洋风味；儿童饮品选择不含咖啡因版本。",
        "afternoon": [
            "午餐后前往陵水，途中让孩子在车上休息。",
            "15:00后入住清水湾，两晚住同一酒店，不再拆分。",
        ],
        "evening": [
            "天气稳定时在酒店开放海滩短时间玩沙。",
            "晚餐尝试陵水酸粉、椰子鸡或酒店附近海南菜。",
        ],
        "kid": "植物园准备驱蚊、防晒和替换衣物；咖啡体验不让儿童误饮含咖啡因饮料。",
        "rain": "小雨可带雨衣游植物园；雷电或暴雨直接取消，提前去陵水入住并使用酒店设施。",
        "booking": "植物园官方常规时间为08:30–17:00，出发前一天再次核对。",
        "nav": [
            ("兴隆热带植物园", "兴隆热带植物园", "万宁"),
            ("清水湾", "陵水清水湾", "陵水"),
        ],
    },
    {
        "day": "D4",
        "route": "陵水完整日",
        "sleep": "陵水清水湾",
        "drive": "根据方案约30–70公里；总驾驶约1–1.5小时",
        "focus": "天气决策日：南湾猴岛 / 海洋欢乐世界 / 酒店避险",
        "morning": [
            "晴朗、风浪小：08:00后前往南湾猴岛，体验跨海索道和猕猴生态。",
            "猴区不携带外露食物和塑料袋，不逗弄、不触摸、不追逐猕猴。",
        ],
        "lunch": "新村镇或正规餐厅用餐；疍家海鲜先确认计价单位、加工费和总价。",
        "afternoon": [
            "午后回酒店午休，17:00后视天气短时玩沙。",
            "若更偏向乐园，可把全天替换为海南海洋欢乐世界，提前核对暑期运营时间。",
        ],
        "evening": [
            "品尝陵水酸粉、气鼓鱼粥、糟粕醋或清补凉。",
            "整理第二天去三亚的行李，避免早晨仓促。",
        ],
        "kid": "猴区由一位成人专门看一个孩子；不让孩子拿饮料、食物或鲜艳袋子靠近猴群。",
        "rain": "小雨可选择海洋欢乐世界的场馆；雷电、暴雨或台风预警时取消所有涉海/户外活动，留在酒店。",
        "booking": "猴岛官方常规时间约08:00–17:20；儿童优惠与身高/年龄相关，携带证件并当日核实。",
        "nav": [
            ("南湾猴岛", "南湾猴岛生态旅游区", "陵水"),
            ("海南海洋欢乐世界", "海南海洋欢乐世界度假区", "陵水"),
            ("陵水清水湾", "陵水清水湾", "陵水"),
        ],
    },
    {
        "day": "D5",
        "route": "陵水 → 三亚",
        "sleep": "三亚亚龙湾",
        "drive": "约55–80公里；纯驾驶约1–1.5小时",
        "focus": "慢退房、转场、三亚度假酒店",
        "morning": [
            "睡到自然醒，早餐后只安排酒店泳池或短时沙滩。",
            "11:00左右退房，确认房间、儿童用品和充电设备无遗漏。",
        ],
        "lunch": "可在清水湾或海棠湾途中用餐，避开边吃边赶路。",
        "afternoon": [
            "前往亚龙湾，入住连续三晚的亲子度假酒店。",
            "午休后熟悉儿童俱乐部、泳池、医务室、沙滩旗和紧急出口。",
        ],
        "evening": [
            "只安排酒店晚餐、泳池或沙滩散步。",
            "不在转场日增加蜈支洲岛、森林公园等大景点。",
        ],
        "kid": "优先确认床围、拼床、儿童早餐、救生衣和儿童活动预约。",
        "rain": "直接入住酒店；使用室内儿童区、商场或酒店活动。",
        "booking": "三亚三晚不换酒店；若重视海滩，优先亚龙湾；若更重视大型亲子设施，可比较海棠湾。",
        "nav": [
            ("亚龙湾", "三亚亚龙湾国家旅游度假区", "三亚"),
            ("海棠湾", "三亚海棠湾", "三亚"),
        ],
    },
    {
        "day": "D6",
        "route": "三亚亲子室内主活动",
        "sleep": "三亚亚龙湾",
        "drive": "亚龙湾往返海棠湾约60–90公里；约1.5–2小时",
        "focus": "亚特兰蒂斯水族馆 + 午休 + 酒店",
        "morning": [
            "早餐后前往失落的空间水族馆，建议10:00左右入场。",
            "控制在2–3小时，不同时购买水世界和多个演出套餐。",
        ],
        "lunch": "在海棠湾商场或度假区用简餐，避开正午继续户外。",
        "afternoon": [
            "回酒店午休，孩子睡醒后使用酒店泳池或儿童俱乐部。",
            "如不想跨湾驾驶，可将本日改为完整酒店日。",
        ],
        "evening": [
            "在亚龙湾周边吃椰子鸡、海南菜或酒店晚餐。",
            "提前查看D7天气，决定是否安排城市日落活动。",
        ],
        "kid": "大型场馆提前约定走失集合点，给孩子佩戴写有家长电话的联系卡。",
        "rain": "水族馆适合普通降雨；若台风导致道路或场馆关闭，则完全留在酒店。",
        "booking": "营业时间和票价变化较多，只通过官方渠道或可靠平台确认；不引用低价非官方票。",
        "nav": [
            ("失落的空间水族馆", "三亚亚特兰蒂斯失落的空间水族馆", "三亚"),
            ("亚龙湾酒店区", "三亚亚龙湾中心广场", "三亚"),
        ],
    },
    {
        "day": "D7",
        "route": "三亚度假与日落",
        "sleep": "三亚亚龙湾",
        "drive": "根据晚间活动约40–70公里；约1–1.5小时",
        "focus": "完整度假日 + 鹿回头或三亚湾日落 + 美食",
        "morning": [
            "酒店早餐、儿童俱乐部、泳池和玩沙，不安排远距离景区。",
            "海况允许时只在有救生员、酒店明确开放的区域亲水。",
        ],
        "lunch": "酒店附近解决，随后完整午休。",
        "afternoon": [
            "16:30后再决定是否前往鹿回头或三亚湾看日落。",
            "如果全家疲劳，直接取消城市活动，继续酒店度假。",
        ],
        "evening": [
            "尝试海南海鲜、港门粉、椰子鸡或糟粕醋。",
            "海鲜选择明码标价餐厅，结账前核对重量、单价和加工费。",
        ],
        "kid": "最后一个完整日不透支体力；回酒店后完成大部分行李打包。",
        "rain": "改为商场、酒店室内活动或提前整理行李；雷电时不在海滩等待日落。",
        "booking": "鹿回头等景区的开放和预约以当天官方信息为准。",
        "nav": [
            ("鹿回头风景区", "三亚鹿回头风景区", "三亚"),
            ("三亚湾椰梦长廊", "三亚湾椰梦长廊", "三亚"),
            ("第一市场", "三亚第一市场", "三亚"),
        ],
    },
    {
        "day": "D8",
        "route": "三亚返程",
        "sleep": "返程",
        "drive": "亚龙湾至凤凰机场约30–40公里；至少预留60分钟",
        "focus": "慢早餐、还车、航班缓冲",
        "morning": [
            "早餐后检查证件、药品、充电器、儿童玩具和湿衣物。",
            "若航班较晚，可继续使用酒店设施，但不新增不可取消景点。",
        ],
        "lunch": "根据航班在酒店或机场附近用餐，避免返程前吃过量生冷海鲜。",
        "afternoon": [
            "按航班起飞前至少3小时抵达还车网点。",
            "还车前加油/充电，拍摄车辆、油量、里程和交车单。",
        ],
        "evening": [
            "保留租车结算、停车、加油和行李托运凭证。",
            "异地还车完成后确认押金与违章查询周期。",
        ],
        "kid": "为机场等待准备一套干衣、零食、饮水和安静玩具。",
        "rain": "台风或航班异常时不要急于退房和还车，先联系航司、酒店和租车公司确认变更政策。",
        "booking": "核对凤凰机场还车网点、摆渡时间、异地还车费和夜间服务费。",
        "nav": [
            ("三亚凤凰国际机场", "三亚凤凰国际机场", "三亚"),
        ],
    },
]


LODGING = [
    {
        "city": "海口",
        "nights": 1,
        "room_nights": 2,
        "area": "骑楼老街 / 国贸",
        "level": "舒适过渡型",
        "budget": "350–450元/间/晚",
        "criteria": "停车方便、早餐、可晚到、距主干道不远",
    },
    {
        "city": "万宁",
        "nights": 1,
        "room_nights": 2,
        "area": "石梅湾 / 兴隆",
        "level": "海湾舒适型",
        "budget": "550–700元/间/晚",
        "criteria": "儿童泳池、室内儿童区、靠近兴隆植物园",
    },
    {
        "city": "陵水",
        "nights": 2,
        "room_nights": 4,
        "area": "清水湾",
        "level": "亲子舒适型",
        "budget": "450–600元/间/晚",
        "criteria": "连续住两晚、停车、泳池、室内雨天活动",
    },
    {
        "city": "三亚",
        "nights": 3,
        "room_nights": 6,
        "area": "亚龙湾",
        "level": "中档亲子度假型",
        "budget": "600–800元/间/晚",
        "criteria": "连续住三晚、儿童俱乐部、指定海滩、早餐政策",
    },
]


BUDGET_ROWS = [
    ("机票", "五人往返海南机票", 5, 1000, 1300, 6000, "目标人均往返约1200元；超出部分从酒店或门票调剂"),
    ("交通", "宽敞五座SUV租金", 8, 250, 350, 2400, "8天；优先后排较宽、后备箱较大的车型"),
    ("交通", "基础保险/补充保障", 8, 40, 60, 400, "确认轮胎、玻璃、底盘和免赔额"),
    ("交通", "海口取车三亚还车费", 1, 200, 400, 300, "以合同为准"),
    ("交通", "儿童安全座椅/增高垫", 16, 20, 40, 480, "2套×8天；必须现场试装后排"),
    ("交通", "燃油", 1, 600, 800, 650, "预计总里程约450–600公里"),
    ("交通", "停车", 1, 200, 350, 250, "酒店和景区停车"),
    ("住宿", "海口2间房×1晚", 2, 350, 450, 700, "舒适过渡型"),
    ("住宿", "万宁2间房×1晚", 2, 550, 700, 1200, "海湾舒适型"),
    ("住宿", "陵水2间房×2晚", 4, 450, 600, 2000, "清水湾亲子舒适型"),
    ("住宿", "三亚2间房×3晚", 6, 600, 800, 3900, "亚龙湾中档亲子度假型"),
    ("餐饮", "五人家庭餐饮", 8, 450, 600, 4000, "含正餐、饮水、儿童零食"),
    ("门票", "南湾猴岛家庭预算", 1, 500, 650, 550, "按3成人+2儿童估算，按身高年龄复核"),
    ("门票", "亚特兰蒂斯水族馆家庭预算", 1, 900, 1300, 1050, "优先家庭套票或可靠平台优惠"),
    ("保障", "旅行意外险", 5, 40, 80, 250, "确认涉水项目和儿童保障"),
    ("保障", "药品、防晒和防蚊", 1, 200, 400, 250, "出发前准备"),
    ("其他", "机动预算", 1, 300, 700, 620, "用于价格波动；不单列大额购物"),
]


CHECKLIST = [
    ("证件", "3位成人身份证、驾驶证"),
    ("证件", "2名儿童身份证/户口簿或可接受的乘机证件"),
    ("证件", "机票、酒店、租车、门票订单离线截图"),
    ("租车", "宽敞五座SUV、后排宽度和行李空间确认"),
    ("租车", "两套儿童安全座椅/增高垫确认"),
    ("租车", "试装后第三位成人可以正常坐下并独立系安全带"),
    ("租车", "异地还车费、保险、押金、里程限制确认"),
    ("租车", "提车360度视频、轮胎、雨刷、刹车、油量"),
    ("健康", "儿童常用退热、过敏、肠胃和晕车药"),
    ("健康", "成人常用药、创可贴、体温计"),
    ("防护", "SPF50+防晒、遮阳帽、太阳镜、防晒衣"),
    ("防护", "驱蚊液、雨衣、折叠伞、防水袋"),
    ("亲水", "儿童救生衣或确认酒店可提供"),
    ("亲水", "泳衣、沙滩鞋、浴巾、换洗衣袋"),
    ("儿童", "水杯、零食、湿巾、安静玩具、联系卡"),
    ("电子", "充电器、充电宝、车载充电线"),
    ("天气", "出发前72小时复核台风、海浪和航班"),
    ("返程", "凤凰机场还车网点和摆渡时间"),
]


SOURCES = [
    (
        "2026年下半年海南热带气旋与强降雨风险",
        "http://www.ce.cn/xwzx/gnsz/gdxw/202607/t20260715_3088395.shtml",
    ),
    (
        "海南暑期涉海和雷雨安全提示",
        "http://www.qxkp.net/tjzx/202508/t20250807_7261073.html",
    ),
    ("海口骑楼老街运营方", "http://www.qiloulaojie.com/"),
    ("兴隆热带植物园官网", "http://www.xlrdzwy.com/"),
    ("南湾猴岛官网", "http://www.monkeyisland.com.cn/cn/home"),
    ("海南海洋欢乐世界官网", "https://www.opresort.cn/web/flhy/index.html"),
    ("三亚亚特兰蒂斯官网", "https://www.atlantissanya.cn/"),
    (
        "海南省交通运输厅2026租车提示（微信文章）",
        "https://mp.weixin.qq.com/s/zvRLSowwASpPIRJWAIzVcQ",
    ),
]


EMERGENCY = [
    ("报警", "110", "人身与治安紧急情况"),
    ("急救", "120", "医疗紧急情况"),
    ("交通事故", "122", "交通事故报警"),
    ("交通运输服务", "12328", "租车/交通服务投诉咨询"),
    ("政务与旅游协助", "12345", "综合投诉与协助"),
    ("海口医院导航", amap("海南省人民医院", "海口"), "按实际位置选择最近院区"),
    ("万宁医院导航", amap("万宁市人民医院", "万宁"), "仅作快速导航入口"),
    ("陵水医院导航", amap("陵水黎族自治县人民医院", "陵水"), "仅作快速导航入口"),
    ("三亚医院导航", amap("三亚中心医院", "三亚"), "仅作快速导航入口"),
]


def nav_links(points: list[tuple[str, str, str]]) -> str:
    return "".join(
        (
            f'<a class="map-link" href="{safe(amap(keyword, city))}" '
            f'target="_blank" rel="noreferrer">导航：{safe(label)}</a>'
        )
        for label, keyword, city in points
    )


def render_day(day: dict[str, object]) -> str:
    return f"""
    <article class="day-card" id="{safe(day['day'])}">
      <div class="day-heading">
        <div class="day-number">{safe(day['day'])}</div>
        <div>
          <h3>{safe(day['route'])}</h3>
          <p class="focus">{safe(day['focus'])}</p>
        </div>
        <div class="sleep-pill">住：{safe(day['sleep'])}</div>
      </div>
      <div class="drive-line">🚗 {safe(day['drive'])}</div>
      <div class="schedule-grid">
        <section>
          <h4>上午</h4>
          <ul>{list_html(day['morning'])}</ul>
        </section>
        <section>
          <h4>午餐 / 午休</h4>
          <p>{safe(day['lunch'])}</p>
        </section>
        <section>
          <h4>下午</h4>
          <ul>{list_html(day['afternoon'])}</ul>
        </section>
        <section>
          <h4>傍晚 / 晚餐</h4>
          <ul>{list_html(day['evening'])}</ul>
        </section>
      </div>
      <div class="notes-grid">
        <div class="note kid"><strong>儿童提醒</strong><span>{safe(day['kid'])}</span></div>
        <div class="note rain"><strong>雨天 / 台风备选</strong><span>{safe(day['rain'])}</span></div>
        <div class="note booking"><strong>预订与复核</strong><span>{safe(day['booking'])}</span></div>
      </div>
      <div class="map-links">{nav_links(day['nav'])}</div>
    </article>
    """


def render_lodging() -> str:
    rows = []
    for item in LODGING:
        rows.append(
            "<tr>"
            f"<td>{safe(item['city'])}</td>"
            f"<td>{safe(item['nights'])}晚 / {safe(item['room_nights'])}间夜</td>"
            f"<td>{safe(item['area'])}</td>"
            f"<td>{safe(item['level'])}</td>"
            f"<td>{safe(item['budget'])}</td>"
            f"<td>{safe(item['criteria'])}</td>"
            "</tr>"
        )
    return "".join(rows)


def render_budget() -> str:
    rows = []
    low_total = 0
    high_total = 0
    target_total = 0
    for category, item, quantity, low, high, target, note in BUDGET_ROWS:
        subtotal_low = quantity * low
        subtotal_high = quantity * high
        low_total += subtotal_low
        high_total += subtotal_high
        target_total += target
        rows.append(
            "<tr>"
            f"<td>{safe(category)}</td>"
            f"<td>{safe(item)}</td>"
            f"<td>{safe(quantity)}</td>"
            f"<td>¥{low:,}–{high:,}</td>"
            f"<td>¥{subtotal_low:,}–{subtotal_high:,}</td>"
            f"<td>¥{target:,}</td>"
            f"<td>{safe(note)}</td>"
            "</tr>"
        )
    if target_total != TARGET_BUDGET:
        raise ValueError(
            f"预算目标分项合计为{target_total}，应为{TARGET_BUDGET}"
        )
    rows.append(
        '<tr class="total-row">'
        '<td colspan="4">规划总预算（包含五人往返机票）</td>'
        f"<td>¥{low_total:,}–{high_total:,}</td>"
        f"<td>目标 ¥{target_total:,}</td>"
        "<td>机票或酒店超预算时，优先取消可选收费项目</td>"
        "</tr>"
    )
    return "".join(rows)


def render_checklist() -> str:
    items = []
    for index, (category, item) in enumerate(CHECKLIST, start=1):
        key = f"check-{index}"
        items.append(
            '<label class="check-item">'
            f'<input type="checkbox" data-check-key="{key}">'
            f"<span><b>{safe(category)}</b>{safe(item)}</span>"
            "</label>"
        )
    return "".join(items)


def render_sources() -> str:
    return "".join(
        (
            f'<li><a href="{safe(url)}" target="_blank" rel="noreferrer">'
            f"{safe(label)}</a></li>"
        )
        for label, url in SOURCES
    )


def render_emergency() -> str:
    rows = []
    for label, value, note in EMERGENCY:
        if value.startswith("http"):
            shown = f'<a href="{safe(value)}" target="_blank">打开地图</a>'
        else:
            shown = f'<a href="tel:{safe(value)}">{safe(value)}</a>'
        rows.append(
            f"<tr><td>{safe(label)}</td><td>{shown}</td><td>{safe(note)}</td></tr>"
        )
    return "".join(rows)


HTML_TEMPLATE = r"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <meta name="color-scheme" content="light">
  <title>@@TITLE@@</title>
  <style>
    :root {
      --ink: #17324d;
      --muted: #60758a;
      --ocean: #0b7895;
      --ocean-soft: #dff5f7;
      --leaf: #18836b;
      --sand: #fff4d8;
      --sun: #f4a63b;
      --coral: #e66a5c;
      --cloud: #f5f8fb;
      --line: #d7e3eb;
      --white: #fff;
      --shadow: 0 14px 40px rgba(23, 50, 77, .12);
    }
    * { box-sizing: border-box; }
    html { scroll-behavior: smooth; }
    body {
      margin: 0;
      color: var(--ink);
      font-family: "Microsoft YaHei", "Noto Sans SC", "PingFang SC", Arial, sans-serif;
      background:
        radial-gradient(circle at 10% 0%, rgba(78, 201, 208, .22), transparent 28rem),
        linear-gradient(180deg, #eefbfc 0, #f8fbfd 36rem, #f6f7f8 100%);
      line-height: 1.65;
    }
    a { color: var(--ocean); }
    button { font: inherit; }
    .hero {
      position: relative;
      overflow: hidden;
      padding: 54px 22px 46px;
      color: white;
      background:
        linear-gradient(125deg, rgba(3, 80, 109, .96), rgba(9, 145, 153, .88)),
        linear-gradient(180deg, #087f9a, #0c8f8c);
    }
    .hero::before, .hero::after {
      content: "";
      position: absolute;
      border-radius: 50%;
      background: rgba(255,255,255,.10);
    }
    .hero::before { width: 340px; height: 340px; right: -90px; top: -180px; }
    .hero::after { width: 220px; height: 220px; left: -100px; bottom: -150px; }
    .hero-inner { max-width: 1120px; margin: auto; position: relative; z-index: 1; }
    .eyebrow {
      display: inline-block;
      padding: 6px 12px;
      border: 1px solid rgba(255,255,255,.45);
      border-radius: 999px;
      font-size: 13px;
      letter-spacing: .08em;
      background: rgba(255,255,255,.10);
    }
    h1 { margin: 14px 0 5px; font-size: clamp(34px, 6vw, 64px); line-height: 1.1; }
    .subtitle { font-size: clamp(18px, 3vw, 25px); margin: 0; opacity: .95; }
    .hero-meta { display: flex; flex-wrap: wrap; gap: 10px; margin-top: 24px; }
    .hero-meta span {
      padding: 8px 12px;
      border-radius: 12px;
      background: rgba(255,255,255,.14);
    }
    .toolbar { display: flex; flex-wrap: wrap; gap: 10px; margin-top: 22px; }
    .toolbar button, .toolbar a {
      border: 0;
      border-radius: 12px;
      padding: 10px 15px;
      cursor: pointer;
      color: var(--ink);
      text-decoration: none;
      background: white;
      box-shadow: 0 8px 20px rgba(0,0,0,.12);
    }
    .toolbar .secondary { background: rgba(255,255,255,.18); color: white; border: 1px solid rgba(255,255,255,.45); }
    .top-nav {
      position: sticky;
      top: 0;
      z-index: 20;
      display: flex;
      gap: 8px;
      padding: 10px 16px;
      overflow-x: auto;
      background: rgba(255,255,255,.94);
      border-bottom: 1px solid var(--line);
      backdrop-filter: blur(10px);
    }
    .top-nav a {
      flex: 0 0 auto;
      padding: 7px 11px;
      border-radius: 999px;
      text-decoration: none;
      color: var(--ink);
      background: var(--cloud);
      font-size: 14px;
    }
    main { max-width: 1120px; margin: auto; padding: 28px 18px 70px; }
    .alert {
      display: grid;
      grid-template-columns: auto 1fr;
      gap: 14px;
      padding: 18px;
      border-radius: 18px;
      background: #fff1ea;
      border: 1px solid #f2b69d;
      box-shadow: var(--shadow);
    }
    .alert-icon { font-size: 28px; }
    .alert strong { color: #9b3e2f; }
    .section { margin-top: 38px; scroll-margin-top: 70px; }
    .section-title { display: flex; align-items: end; gap: 12px; margin-bottom: 16px; }
    .section-title h2 { margin: 0; font-size: 28px; }
    .section-title span { color: var(--muted); font-size: 14px; }
    .route {
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      gap: 12px;
    }
    .route-stop {
      position: relative;
      min-height: 150px;
      padding: 18px;
      border-radius: 20px;
      background: white;
      box-shadow: var(--shadow);
      border-top: 5px solid var(--ocean);
    }
    .route-stop:nth-child(2) { border-color: var(--sun); }
    .route-stop:nth-child(3) { border-color: var(--leaf); }
    .route-stop:nth-child(4) { border-color: var(--coral); }
    .route-stop h3 { margin: 0; font-size: 24px; }
    .route-stop .nights { font-size: 32px; font-weight: 800; color: var(--ocean); }
    .route-stop p { margin: 4px 0 0; color: var(--muted); }
    .principles {
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 12px;
    }
    .principle {
      background: white;
      padding: 18px;
      border-radius: 18px;
      border: 1px solid var(--line);
    }
    .principle b { display: block; margin-bottom: 5px; color: var(--ocean); }
    .day-card {
      margin: 18px 0;
      padding: 22px;
      border-radius: 24px;
      background: rgba(255,255,255,.98);
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
      scroll-margin-top: 70px;
    }
    .day-heading {
      display: grid;
      grid-template-columns: auto 1fr auto;
      gap: 15px;
      align-items: center;
    }
    .day-number {
      display: grid;
      place-items: center;
      width: 58px;
      height: 58px;
      border-radius: 18px;
      color: white;
      background: linear-gradient(145deg, var(--ocean), var(--leaf));
      font-size: 20px;
      font-weight: 800;
    }
    .day-heading h3 { margin: 0; font-size: 24px; }
    .focus { margin: 2px 0 0; color: var(--muted); }
    .sleep-pill {
      padding: 8px 11px;
      border-radius: 12px;
      background: var(--sand);
      font-size: 14px;
      font-weight: 700;
    }
    .drive-line {
      margin: 16px 0;
      padding: 10px 13px;
      border-radius: 12px;
      background: var(--cloud);
      color: var(--muted);
    }
    .schedule-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 12px;
    }
    .schedule-grid section {
      padding: 15px;
      border-radius: 16px;
      background: #f7fbfc;
      border: 1px solid #dcebed;
    }
    .schedule-grid h4 { margin: 0 0 6px; color: var(--ocean); }
    .schedule-grid p, .schedule-grid ul { margin: 0; }
    .schedule-grid ul { padding-left: 20px; }
    .notes-grid {
      display: grid;
      grid-template-columns: repeat(3, minmax(0, 1fr));
      gap: 10px;
      margin-top: 12px;
    }
    .note { padding: 13px; border-radius: 14px; font-size: 14px; }
    .note strong { display: block; margin-bottom: 4px; }
    .note span { color: #40576a; }
    .note.kid { background: #fff4d8; }
    .note.rain { background: #e9f1ff; }
    .note.booking { background: #e9f8ee; }
    .map-links { display: flex; flex-wrap: wrap; gap: 8px; margin-top: 13px; }
    .map-link {
      display: inline-block;
      padding: 7px 11px;
      border-radius: 999px;
      text-decoration: none;
      color: white;
      background: var(--ocean);
      font-size: 13px;
    }
    .table-wrap { overflow-x: auto; border-radius: 18px; box-shadow: var(--shadow); }
    table { width: 100%; border-collapse: collapse; background: white; font-size: 14px; }
    th, td { padding: 11px 12px; border-bottom: 1px solid var(--line); text-align: left; vertical-align: top; }
    th { background: #e7f5f6; color: #17566a; white-space: nowrap; }
    .total-row td { font-weight: 800; background: var(--sand); }
    .food-grid, .safety-grid {
      display: grid;
      grid-template-columns: repeat(4, 1fr);
      gap: 12px;
    }
    .food-card, .safety-card {
      padding: 16px;
      background: white;
      border: 1px solid var(--line);
      border-radius: 18px;
    }
    .food-card h3, .safety-card h3 { margin: 0 0 6px; color: var(--ocean); }
    .food-card p, .safety-card p { margin: 0; color: var(--muted); }
    .checklist {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 8px;
    }
    .check-item {
      display: flex;
      gap: 10px;
      align-items: start;
      padding: 11px 12px;
      background: white;
      border: 1px solid var(--line);
      border-radius: 12px;
      cursor: pointer;
    }
    .check-item input { margin-top: 5px; transform: scale(1.15); }
    .check-item b {
      display: inline-block;
      min-width: 54px;
      margin-right: 8px;
      color: var(--ocean);
    }
    .check-item.done { opacity: .55; text-decoration: line-through; }
    .source-box {
      padding: 18px;
      border-radius: 18px;
      background: white;
      border: 1px solid var(--line);
    }
    .source-box li { margin: 7px 0; }
    .footer {
      margin-top: 36px;
      padding: 24px;
      border-radius: 20px;
      background: var(--ink);
      color: white;
    }
    .footer p { margin: 5px 0; opacity: .88; }
    @media (max-width: 850px) {
      .route, .principles, .food-grid, .safety-grid { grid-template-columns: repeat(2, 1fr); }
      .notes-grid { grid-template-columns: 1fr; }
      .day-heading { grid-template-columns: auto 1fr; }
      .sleep-pill { grid-column: 1 / -1; }
    }
    @media (max-width: 600px) {
      .hero { padding-top: 38px; }
      main { padding-left: 12px; padding-right: 12px; }
      .route, .principles, .food-grid, .safety-grid, .schedule-grid, .checklist { grid-template-columns: 1fr; }
      .day-card { padding: 16px; border-radius: 18px; }
      .day-heading h3 { font-size: 20px; }
      .alert { grid-template-columns: 1fr; }
    }
    @page { size: A4; margin: 11mm; }
    @media print {
      body { background: white; font-size: 10pt; }
      .hero {
        padding: 25px 20px;
        -webkit-print-color-adjust: exact;
        print-color-adjust: exact;
      }
      .hero h1 { font-size: 30pt; }
      .top-nav, .toolbar { display: none !important; }
      main { max-width: none; padding: 10px 0; }
      .alert, .route-stop, .day-card, .table-wrap { box-shadow: none; }
      .day-card { break-before: page; break-inside: avoid; }
      .day-card:first-of-type { break-before: auto; }
      .section { break-before: page; margin-top: 0; }
      #route, #days { break-before: auto; }
      .route, .principles, .food-grid, .safety-grid { grid-template-columns: repeat(2, 1fr); }
      .schedule-grid { grid-template-columns: repeat(2, 1fr); }
      .notes-grid { grid-template-columns: 1fr; }
      a { color: var(--ocean); text-decoration: none; }
      .check-item input { appearance: none; width: 13px; height: 13px; border: 1px solid #555; }
      .footer { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
    }
  </style>
</head>
<body>
  <header class="hero">
    <div class="hero-inner">
      <span class="eyebrow">3位成人 · 2名5–6岁儿童 · 宽敞五座SUV</span>
      <h1>@@TITLE@@</h1>
      <p class="subtitle">@@SUBTITLE@@</p>
      <div class="hero-meta">
        <span>📅 @@TRAVEL_WINDOW@@</span>
        <span>🏨 1 + 1 + 2 + 3晚</span>
        <span>🚗 约450–600公里</span>
        <span>💰 含机票目标约¥25,000</span>
        <span>🌦 8月台风季双方案</span>
      </div>
      <div class="toolbar">
        <button type="button" onclick="window.print()">打印 / 另存PDF</button>
        <a href="#days">查看逐日安排</a>
        <button class="secondary" type="button" onclick="resetChecks()">清空已勾选清单</button>
      </div>
    </div>
  </header>

  <nav class="top-nav">
    <a href="#route">路线</a>
    <a href="#days">逐日行程</a>
    <a href="#hotels">住宿</a>
    <a href="#food">美食</a>
    <a href="#budget">预算</a>
    <a href="#packing">清单</a>
    <a href="#emergency">应急</a>
    <a href="#sources">资料来源</a>
  </nav>

  <main>
    <section class="alert">
      <div class="alert-icon">⚠️</div>
      <div>
        <strong>8月是高温、雷雨和台风高发期。</strong>
        本攻略不把轮渡、海岛和海上项目设成不可替换核心行程。出发前72小时及每天早晨查看官方气象、海浪、道路和景区公告；出现雷电、暴雨或台风预警时，立即取消海滩、索道、乘船和山区活动。
      </div>
    </section>

    <section class="section" id="route">
      <div class="section-title"><h2>路线与住宿</h2><span>线性东线，不走回头路</span></div>
      <div class="route">
        <article class="route-stop"><h3>海口</h3><div class="nights">1晚</div><p>取车 · 骑楼 · 过渡</p></article>
        <article class="route-stop"><h3>万宁</h3><div class="nights">1晚</div><p>石梅湾 · 兴隆植物园</p></article>
        <article class="route-stop"><h3>陵水</h3><div class="nights">2晚</div><p>清水湾 · 猴岛 / 海洋乐园</p></article>
        <article class="route-stop"><h3>三亚</h3><div class="nights">3晚</div><p>亚龙湾 · 水族馆 · 度假</p></article>
      </div>
      <div class="principles" style="margin-top:12px">
        <div class="principle"><b>每天一个主活动</b>不以景点数量衡量旅行质量。</div>
        <div class="principle"><b>午间固定休息</b>11:30–15:30尽量室内或午睡。</div>
        <div class="principle"><b>天气当天决策</b>海边、猴岛与乐园均准备替代方案。</div>
      </div>
      <div class="alert" style="margin-top:12px;background:#fff4d8;border-color:#e5c676">
        <div class="alert-icon">🚙</div>
        <div><strong>五座车条件：</strong>两套儿童约束装置安装后，后排中间成人必须能正常坐下并独立系安全带，后备箱能装下全部行李；提车实测不满足时应升级七座，不能牺牲安全换预算。</div>
      </div>
    </section>

    <section class="section" id="days">
      <div class="section-title"><h2>8天逐日安排</h2><span>时间均为建议窗口，可随航班调整</span></div>
      @@DAY_CARDS@@
    </section>

    <section class="section" id="hotels">
      <div class="section-title"><h2>住宿选择</h2><span>两间房，共14间夜</span></div>
      <div class="table-wrap">
        <table>
          <thead><tr><th>城市</th><th>晚数</th><th>推荐区域</th><th>类型</th><th>规划价</th><th>筛选条件</th></tr></thead>
          <tbody>@@LODGING_ROWS@@</tbody>
        </table>
      </div>
      <p>订房时重点确认：两间房是否同层/相邻、儿童早餐身高政策、床围或拼床、泳池开放、室内儿童区、停车费和台风取消政策。不要在三亚三晚之间再次换酒店。</p>
    </section>

    <section class="section" id="food">
      <div class="section-title"><h2>沿途美食</h2><span>儿童优先清淡，成人再尝鲜</span></div>
      <div class="food-grid">
        <article class="food-card"><h3>海口</h3><p>辣汤饭、海南粉、老爸茶、清补凉。首晚避免过量生冷。</p></article>
        <article class="food-card"><h3>万宁</h3><p>后安粉、兴隆咖啡、南洋菜、东山羊、和乐蟹。</p></article>
        <article class="food-card"><h3>陵水</h3><p>陵水酸粉、气鼓鱼粥、疍家海鲜、糟粕醋、炒冰。</p></article>
        <article class="food-card"><h3>三亚</h3><p>椰子鸡、港门粉、海南菜、明码标价海鲜和热带水果。</p></article>
      </div>
      <p>海鲜消费原则：确认品种、单价、计价单位、重量、加工费和总价；孩子不吃来源不明的生腌、生食或放置过久海鲜。</p>
    </section>

    <section class="section" id="budget">
      <div class="section-title"><h2>预算控制</h2><span>包含五人往返机票，目标约¥25,000</span></div>
      <div class="table-wrap">
        <table>
          <thead><tr><th>类别</th><th>项目</th><th>数量</th><th>单价区间</th><th>市场区间小计</th><th>目标额度</th><th>说明</th></tr></thead>
          <tbody>@@BUDGET_ROWS@@</tbody>
        </table>
      </div>
      <p><strong>控制线：</strong>五人往返机票尽量控制在¥6,000，14个酒店间夜控制在¥7,800，车辆相关约¥4,480，餐饮约¥4,000，门票约¥1,600，其余约¥1,120。若机票超过¥6,000，优先下调三亚房价、取消亚特兰蒂斯收费项目或减少餐饮支出。</p>
      <p>海南高速通常不另设收费站；汽油车通行附加费体现在油价中。实际租车、保险、儿童座椅和异地还车费用以合同为准。</p>
    </section>

    <section class="section" id="packing">
      <div class="section-title"><h2>出发前清单</h2><span>勾选状态只保存在当前浏览器</span></div>
      <div class="checklist">@@CHECKLIST@@</div>
    </section>

    <section class="section" id="emergency">
      <div class="section-title"><h2>应急与安全</h2><span>危险时先确保人身安全，再处理行程</span></div>
      <div class="safety-grid">
        <article class="safety-card"><h3>雷雨</h3><p>立即离开海滩、泳池边、孤立树木和高处；进入有防雷设施的建筑。</p></article>
        <article class="safety-card"><h3>台风</h3><p>暂停涉海、索道和山区活动；确认航班、高铁、道路、景区和酒店安排。</p></article>
        <article class="safety-card"><h3>儿童亲水</h3><p>只在明确开放区域，由成人一对一看护；救生衣不能替代看护。</p></article>
        <article class="safety-card"><h3>车辆事故</h3><p>先报警和联系租车公司，保留现场照片；未经同意不要自行维修。</p></article>
      </div>
      <div class="table-wrap" style="margin-top:12px">
        <table>
          <thead><tr><th>项目</th><th>电话 / 导航</th><th>说明</th></tr></thead>
          <tbody>@@EMERGENCY_ROWS@@</tbody>
        </table>
      </div>
    </section>

    <section class="section" id="sources">
      <div class="section-title"><h2>资料来源与复核</h2><span>整理日期：@@RESEARCHED_AT@@</span></div>
      <div class="source-box">
        <ul>@@SOURCE_ROWS@@</ul>
        <p><strong>说明：</strong>景区开放、儿童票、海况、酒店和租车价格均可能变化。本攻略是行程决策底稿，不替代景区、气象、航司、酒店及租车公司的当天公告。</p>
      </div>
    </section>

    <footer class="footer">
      <strong>最重要的行程原则</strong>
      <p>孩子状态优先于景点；天气安全优先于订单；酒店休息也是正式行程。</p>
      <p>路线：海口1晚 → 万宁1晚 → 陵水2晚 → 三亚3晚。</p>
    </footer>
  </main>

  <script>
    (function () {
      var boxes = document.querySelectorAll('input[data-check-key]');
      boxes.forEach(function (box) {
        var key = 'hainan-trip-' + box.getAttribute('data-check-key');
        try { box.checked = localStorage.getItem(key) === '1'; } catch (e) {}
        box.closest('.check-item').classList.toggle('done', box.checked);
        box.addEventListener('change', function () {
          try { localStorage.setItem(key, box.checked ? '1' : '0'); } catch (e) {}
          box.closest('.check-item').classList.toggle('done', box.checked);
        });
      });
    }());
    function resetChecks() {
      document.querySelectorAll('input[data-check-key]').forEach(function (box) {
        box.checked = false;
        box.closest('.check-item').classList.remove('done');
        try {
          localStorage.removeItem('hainan-trip-' + box.getAttribute('data-check-key'));
        } catch (e) {}
      });
    }
  </script>
</body>
</html>
"""


def build_html() -> None:
    content = (
        HTML_TEMPLATE.replace("@@TITLE@@", safe(TITLE))
        .replace("@@SUBTITLE@@", safe(SUBTITLE))
        .replace("@@TRAVEL_WINDOW@@", safe(TRAVEL_WINDOW))
        .replace("@@RESEARCHED_AT@@", safe(RESEARCHED_AT))
        .replace("@@DAY_CARDS@@", "".join(render_day(day) for day in DAYS))
        .replace("@@LODGING_ROWS@@", render_lodging())
        .replace("@@BUDGET_ROWS@@", render_budget())
        .replace("@@CHECKLIST@@", render_checklist())
        .replace("@@EMERGENCY_ROWS@@", render_emergency())
        .replace("@@SOURCE_ROWS@@", render_sources())
    )
    HTML_PATH.write_text(content, encoding="utf-8")


NAVY = "17324D"
OCEAN = "0B7895"
OCEAN_LIGHT = "DFF5F7"
SAND = "FFF4D8"
CORAL = "E66A5C"
WHITE = "FFFFFF"
LINE = "D7E3EB"


def style_sheet(ws, widths: dict[int, float], freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color=LINE)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=OCEAN)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.auto_filter.ref = ws.dimensions


def build_workbook() -> None:
    wb = Workbook()
    overview = wb.active
    overview.title = "行程总览"
    overview.append(["项目", "内容"])
    overview_rows = [
        ("旅行窗口", TRAVEL_WINDOW),
        ("人数", "3位成人 + 2名5–6岁儿童"),
        ("路线", SUBTITLE),
        ("车型", "空间较大的五座SUV；两套儿童约束装置必须现场试装"),
        ("目标总预算", f"¥{TARGET_BUDGET:,}，包含五人往返机票"),
        ("预计里程", "约450–600公里"),
        ("行程原则", "每天一个主活动；午间休息；恶劣天气取消涉海活动"),
        ("HTML攻略", HTML_PATH.name),
        ("PDF攻略", PDF_PATH.name),
        ("动态信息整理日期", RESEARCHED_AT),
    ]
    for row in overview_rows:
        overview.append(row)
    style_sheet(overview, {1: 24, 2: 92})
    overview.auto_filter.ref = "A1:B1"
    for row in range(2, overview.max_row + 1):
        overview.cell(row, 1).font = Font(bold=True, color=OCEAN)

    daily = wb.create_sheet("逐日行程")
    daily.append(
        [
            "天数",
            "路线",
            "住宿",
            "驾驶",
            "核心安排",
            "上午",
            "午餐/午休",
            "下午",
            "傍晚/晚餐",
            "儿童提醒",
            "雨天备选",
            "预订复核",
            "首要导航",
        ]
    )
    for day in DAYS:
        first_label, first_keyword, first_city = day["nav"][0]
        daily.append(
            [
                day["day"],
                day["route"],
                day["sleep"],
                day["drive"],
                day["focus"],
                "\n".join(f"• {item}" for item in day["morning"]),
                day["lunch"],
                "\n".join(f"• {item}" for item in day["afternoon"]),
                "\n".join(f"• {item}" for item in day["evening"]),
                day["kid"],
                day["rain"],
                day["booking"],
                first_label,
            ]
        )
        link_cell = daily.cell(daily.max_row, 13)
        link_cell.hyperlink = amap(first_keyword, first_city)
        link_cell.style = "Hyperlink"
    style_sheet(
        daily,
        {
            1: 8,
            2: 20,
            3: 24,
            4: 30,
            5: 28,
            6: 42,
            7: 35,
            8: 42,
            9: 42,
            10: 42,
            11: 48,
            12: 44,
            13: 20,
        },
    )
    for row in range(2, daily.max_row + 1):
        daily.row_dimensions[row].height = 110
        daily.cell(row, 1).fill = PatternFill("solid", fgColor=OCEAN_LIGHT)
        daily.cell(row, 1).font = Font(bold=True, color=OCEAN)

    booking = wb.create_sheet("预订与预算")
    booking.append(
        [
            "类别",
            "项目",
            "数量",
            "低位单价",
            "高位单价",
            "低位小计",
            "高位小计",
            "目标额度",
            "实际金额",
            "预算差额",
            "状态",
            "取消/确认期限",
            "订单号",
            "说明",
        ]
    )
    for index, (category, item, quantity, low, high, target, note) in enumerate(
        BUDGET_ROWS, start=2
    ):
        booking.append(
            [
                category,
                item,
                quantity,
                low,
                high,
                f"=C{index}*D{index}",
                f"=C{index}*E{index}",
                target,
                None,
                f'=IF(I{index}="","",H{index}-I{index})',
                "未预订",
                None,
                None,
                note,
            ]
        )
    total_row = booking.max_row + 1
    booking.cell(total_row, 1, "合计")
    booking.cell(total_row, 6, f"=SUM(F2:F{total_row - 1})")
    booking.cell(total_row, 7, f"=SUM(G2:G{total_row - 1})")
    booking.cell(total_row, 8, f"=SUM(H2:H{total_row - 1})")
    booking.cell(total_row, 9, f"=SUM(I2:I{total_row - 1})")
    booking.cell(total_row, 10, f"=H{total_row}-I{total_row}")
    for cell in booking[total_row]:
        cell.fill = PatternFill("solid", fgColor=SAND)
        cell.font = Font(bold=True, color=NAVY)
    status_validation = DataValidation(
        type="list",
        formula1='"未预订,待确认,已预订,已支付,已取消"',
        allow_blank=True,
    )
    booking.add_data_validation(status_validation)
    status_validation.add(f"K2:K{total_row - 1}")
    for row in range(2, total_row + 1):
        for column in range(4, 11):
            booking.cell(row, column).number_format = '¥#,##0.00'
    style_sheet(
        booking,
        {
            1: 12,
            2: 30,
            3: 10,
            4: 14,
            5: 14,
            6: 14,
            7: 14,
            8: 14,
            9: 14,
            10: 14,
            11: 13,
            12: 18,
            13: 20,
            14: 46,
        },
    )
    booking.conditional_formatting.add(
        f"I2:I{total_row - 1}",
        ColorScaleRule(
            start_type="min",
            start_color="E9F8EE",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFF4D8",
            end_type="max",
            end_color="FADBD8",
        ),
    )

    hotels = wb.create_sheet("酒店筛选")
    hotels.append(["城市", "晚数", "间夜", "区域", "档次", "预算", "筛选条件", "候选酒店", "订单状态"])
    for item in LODGING:
        hotels.append(
            [
                item["city"],
                item["nights"],
                item["room_nights"],
                item["area"],
                item["level"],
                item["budget"],
                item["criteria"],
                None,
                "未预订",
            ]
        )
    hotel_status_validation = DataValidation(
        type="list",
        formula1='"未预订,待确认,已预订,已支付,已取消"',
        allow_blank=True,
    )
    hotels.add_data_validation(hotel_status_validation)
    hotel_status_validation.add(f"I2:I{hotels.max_row}")
    style_sheet(
        hotels,
        {1: 12, 2: 9, 3: 9, 4: 22, 5: 22, 6: 22, 7: 52, 8: 30, 9: 13},
    )

    checklist = wb.create_sheet("打包清单")
    checklist.append(["类别", "物品/事项", "状态", "负责人", "备注"])
    for category, item in CHECKLIST:
        checklist.append([category, item, "未完成", None, None])
    check_validation = DataValidation(
        type="list", formula1='"未完成,已完成,不适用"', allow_blank=False
    )
    checklist.add_data_validation(check_validation)
    check_validation.add(f"C2:C{checklist.max_row}")
    style_sheet(checklist, {1: 13, 2: 62, 3: 13, 4: 16, 5: 36})

    links = wb.create_sheet("地图与资料")
    links.append(["类型", "名称", "链接", "说明"])
    seen: set[tuple[str, str]] = set()
    for day in DAYS:
        for label, keyword, city in day["nav"]:
            key = (label, city)
            if key in seen:
                continue
            seen.add(key)
            url = amap(keyword, city)
            links.append(["高德导航", label, url, f"{city}搜索入口"])
            links.cell(links.max_row, 3).hyperlink = url
            links.cell(links.max_row, 3).style = "Hyperlink"
    for label, url in SOURCES:
        links.append(["资料来源", label, url, f"整理日期 {RESEARCHED_AT}"])
        links.cell(links.max_row, 3).hyperlink = url
        links.cell(links.max_row, 3).style = "Hyperlink"
    for label, value, note in EMERGENCY:
        if value.startswith("http"):
            links.append(["应急导航", label, value, note])
            links.cell(links.max_row, 3).hyperlink = value
            links.cell(links.max_row, 3).style = "Hyperlink"
    style_sheet(links, {1: 14, 2: 34, 3: 72, 4: 40})

    emergency = wb.create_sheet("应急卡")
    emergency.append(["项目", "电话/入口", "说明"])
    for label, value, note in EMERGENCY:
        emergency.append([label, value, note])
        if value.startswith("http"):
            emergency.cell(emergency.max_row, 2).hyperlink = value
            emergency.cell(emergency.max_row, 2).style = "Hyperlink"
    emergency.append(
        [
            "台风处理",
            "先联系航司、酒店、租车公司",
            "不要在未确认改签/续住前急于退房还车",
        ]
    )
    emergency.append(
        [
            "海滩安全",
            "只在明确开放区域活动",
            "雷电、风浪或红旗时立即离开海滩",
        ]
    )
    style_sheet(emergency, {1: 22, 2: 45, 3: 70})

    wb.save(XLSX_PATH)


async def build_pdf() -> None:
    edge_candidates = [
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    ]
    executable = next((path for path in edge_candidates if path.exists()), None)
    async with async_playwright() as playwright:
        launch_kwargs: dict[str, object] = {"headless": True}
        if executable:
            launch_kwargs["executable_path"] = str(executable)
        browser = await playwright.chromium.launch(**launch_kwargs)
        page = await browser.new_page(viewport={"width": 1440, "height": 1000})
        await page.goto(HTML_PATH.as_uri(), wait_until="load")
        await page.emulate_media(media="print")
        await page.pdf(
            path=str(PDF_PATH),
            format="A4",
            print_background=True,
            prefer_css_page_size=True,
            margin={"top": "8mm", "right": "8mm", "bottom": "8mm", "left": "8mm"},
        )
        await browser.close()


async def main() -> None:
    build_html()
    build_workbook()
    await build_pdf()
    for path in (HTML_PATH, PDF_PATH, XLSX_PATH):
        print(f"{path.name}: {path.stat().st_size:,} bytes")


if __name__ == "__main__":
    asyncio.run(main())
